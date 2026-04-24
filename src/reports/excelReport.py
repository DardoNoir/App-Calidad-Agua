from pathlib import Path
from typing import List, Optional, cast
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from src.utils.datetime import format_fecha_display

import matplotlib.pyplot as plt

from src.data.repository import (
    traerMedicionPorId,
    traerLugarPorId,
    traerUsuarioPorId,
    traerValoresMedidosDeMedicion,
    traerValoresReferenciaPorParametro
)

from src.charts.datasource import traerSerieCompletaParametroLugar
from src.charts.timeseries import graficarParametro
from src.analysis.evaluacion import ReferenciaParametro


# ==========================================================
# FUNCIONES AUXILIARES DE ESTILO
# ==========================================================

from openpyxl.utils import get_column_letter

def ajustar_ancho_columnas(ws: Worksheet):
    for col_idx, col in enumerate(ws.columns, start=1):
        max_length = 0
        col_letter = get_column_letter(col_idx)

        for cell in col:
            if cell.value is not None:
                max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[col_letter].width = max_length + 2

def aplicar_bordes(ws: Worksheet):
    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for row in ws.iter_rows():
        for cell in row:
            cell.border = border


def estilizar_encabezado(ws: Worksheet):
    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")


# ==========================================================
# FUNCIÓN PRINCIPAL
# ==========================================================

def generarExcelMediciones(
    idsMediciones: List[int],
    rutaGuardar: Path,
):

    wb = Workbook()

    # ======================================================
    # HOJA 1 - RESUMEN
    # ======================================================

    ws_resumen = cast(Worksheet, wb.active)
    ws_resumen.title = "Resumen Mediciones"

    headers = [
        "ID",
        "Fecha",
        "Provincia",
        "Ciudad",
        "Cuerpo de Agua",
        "Tomada Por",
        "ICA",
        "Clasificación",
        "Observaciones"
    ]

    ws_resumen.append(headers)

    for idMedicion in idsMediciones:

        medicion = traerMedicionPorId(idMedicion)
        if medicion is None:
            continue

        lugar = traerLugarPorId(medicion["id_lugar"])
        if lugar is None:
            raise ValueError("No existe el Lugar")

        usuario = traerUsuarioPorId(medicion["medida_tomada_por"])
        if usuario is None:
            raise ValueError("No existe el usuario")

        ws_resumen.append([
            medicion["id"],
            format_fecha_display(medicion["fecha_medicion"]),
            lugar["provincia"],
            lugar["ciudad"],
            lugar["cuerpo_agua"],
            usuario["nombre_usuario"],
            medicion["ica"],
            medicion["clasificacion_ica"],
            medicion["comentarios"]
        ])

    # ---- Estilo resumen ----
    estilizar_encabezado(ws_resumen)
    aplicar_bordes(ws_resumen)
    ajustar_ancho_columnas(ws_resumen)

    # ---- Convertir en tabla real Excel ----
    tabla = Table(
        displayName="TablaResumen",
        ref=f"A1:I{ws_resumen.max_row}"
    )

    estilo_tabla = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False
    )

    tabla.tableStyleInfo = estilo_tabla
    ws_resumen.add_table(tabla)

    # ======================================================
    # HOJAS DETALLADAS
    # ======================================================

    for idMedicion in idsMediciones:

        medicion = traerMedicionPorId(idMedicion)
        if medicion is None:
            continue

        lugar = traerLugarPorId(medicion["id_lugar"])
        if lugar is None:
            raise ValueError("No existe el Lugar")

        usuario = traerUsuarioPorId(medicion["medida_tomada_por"])
        if usuario is None:
            raise ValueError("No existe el usuario")

        valores = traerValoresMedidosDeMedicion(idMedicion)

        # --------------------------------------------------
        # HOJA DATOS
        # --------------------------------------------------

        ws_datos = cast(
            Worksheet,
            wb.create_sheet(f"Datos {idMedicion}")
        )

        ws_datos.append(["Medición N°", medicion["id"]])
        ws_datos.append(["Fecha", format_fecha_display(medicion["fecha_medicion"])])
        ws_datos.append([
            "Lugar",
            f"{lugar['provincia']} - {lugar['ciudad']} - {lugar['cuerpo_agua']}"
        ])
        ws_datos.append(["Tomada por", usuario["nombre_usuario"]])
        ws_datos.append(["ICA", medicion["ica"]])
        ws_datos.append(["Clasificación", medicion["clasificacion_ica"]])
        ws_datos.append([])

        ws_datos.append([
            "Parámetro",
            "Valor",
            "Unidad",
            "Estado",
            "Desviación"
        ])

        for valor in valores:
            ws_datos.append([
                valor["nombre_parametro"],
                valor["valor"],
                valor["unidad"],
                valor["estado"],
                valor["desviacion"]
            ])

        estilizar_encabezado(ws_datos)
        aplicar_bordes(ws_datos)
        ajustar_ancho_columnas(ws_datos)

        # --------------------------------------------------
        # HOJA GRÁFICOS
        # --------------------------------------------------

        ws_graficos = cast(
            Worksheet,
            wb.create_sheet(f"Graficos {idMedicion}")
        )

        fila_imagen = 1

        for valor in valores:

            idParametro = valor["id_parametro"]
            nombreParametro = valor["nombre_parametro"]
            unidad = valor["unidad"]

            serie = traerSerieCompletaParametroLugar(
                medicion["id_lugar"],
                idParametro
            )

            if not serie:
                continue

            fechas = [s["fecha"] for s in serie]
            valores_hist = [float(s["valor"]) for s in serie]

            referencia = traerValoresReferenciaPorParametro(
                idParametro
            )

            ref = cast(ReferenciaParametro | None, referencia)

            fig = graficarParametro(
                fechas=fechas,
                valores=valores_hist,
                nombreParametro=nombreParametro,
                unidad=unidad,
                referencia=ref,
                titulo=f"en {lugar['cuerpo_agua']}"
            )

            img_bytes = BytesIO()
            fig.savefig(img_bytes, format="png")
            img_bytes.seek(0)
            plt.close(fig)

            img = XLImage(img_bytes)
            img.width = 600
            img.height = 350

            ws_graficos.add_image(img, f"A{fila_imagen}")
            fila_imagen += 20

    wb.save(rutaGuardar)
    return rutaGuardar